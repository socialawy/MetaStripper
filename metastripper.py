import os
import sys
import logging
import shutil
import zipfile
from datetime import datetime
import tempfile
from tkinter import Tk, ttk, filedialog, messagebox, BooleanVar, IntVar, Text, END, VERTICAL
from PIL import Image
try:
    import imageio
    IMAGEIO_AVAILABLE = True
except ImportError:
    IMAGEIO_AVAILABLE = False
try:
    import ffmpeg
    FFMPEG_AVAILABLE = True
except ImportError:
    FFMPEG_AVAILABLE = False
from PyPDF2 import PdfReader, PdfWriter
from pptx import Presentation
from docx import Document
from openpyxl import load_workbook
from odf import text, teletype
from odf.opendocument import load as load_odf
import mutagen
from hachoir.parser import createParser
from hachoir.metadata import extractMetadata
import rarfile
import py7zr

class MetaStripper:
    def __init__(self, root):
        self.root = root
        self.root.title("MetaStripper - Remove File Metadata")
        self.root.geometry("700x450")
        self.cleanup_temp()
        self.setup_logging()
        self.setup_ui()

    def setup_logging(self):
        """Initialize logging to file and console."""
        self.logger = logging.getLogger('MetaStripper')
        self.logger.setLevel(logging.DEBUG)
        log_file = os.path.join(tempfile.gettempdir(), 'metastripper.log')
        fh = logging.FileHandler(log_file)
        fh.setLevel(logging.DEBUG)
        ch = logging.StreamHandler(sys.stdout)
        ch.setLevel(logging.INFO)
        formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
        fh.setFormatter(formatter)
        ch.setFormatter(formatter)
        self.logger.addHandler(fh)
        self.logger.addHandler(ch)

    def setup_ui(self):
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill='both', expand=True)

        # File/Folder selection
        ttk.Label(main_frame, text="Select File(s)/Folder:").grid(row=0, column=0, sticky='w')
        self.file_entry = ttk.Entry(main_frame, width=50)
        self.file_entry.grid(row=0, column=1, sticky=('w', 'e'), padx=5)
        ttk.Button(main_frame, text="Browse", command=self.browse_files).grid(row=0, column=2)

        # Output location
        ttk.Label(main_frame, text="Output Folder:").grid(row=1, column=0, sticky='w')
        self.output_entry = ttk.Entry(main_frame, width=50)
        self.output_entry.grid(row=1, column=1, sticky=('w', 'e'), padx=5)
        ttk.Button(main_frame, text="Browse", command=self.browse_output).grid(row=1, column=2)
        self.same_as_input_var = BooleanVar(value=True)
        ttk.Checkbutton(main_frame, text="Same as input folder", variable=self.same_as_input_var, 
                        command=self.toggle_output).grid(row=2, column=1, sticky='w')

        # Options
        options_frame = ttk.LabelFrame(main_frame, text="Cleaning Options", padding="10")
        options_frame.grid(row=3, column=0, columnspan=3, sticky=('w', 'e'), pady=10)

        self.remove_all = BooleanVar(value=True)
        ttk.Checkbutton(options_frame, text="Remove all metadata", variable=self.remove_all).pack(anchor='w')
        self.keep_copyright = BooleanVar()
        ttk.Checkbutton(options_frame, text="Keep copyright info", variable=self.keep_copyright).pack(anchor='w')
        self.keep_date = BooleanVar()
        ttk.Checkbutton(options_frame, text="Keep creation date", variable=self.keep_date).pack(anchor='w')
        self.recursive = BooleanVar()
        ttk.Checkbutton(options_frame, text="Process folders recursively", variable=self.recursive).pack(anchor='w')
        self.backup = BooleanVar()
        ttk.Checkbutton(options_frame, text="Create backup before cleaning", variable=self.backup).pack(anchor='w')
        self.size_limit = IntVar(value=0)
        ttk.Label(options_frame, text="Max file size (MB):").pack(anchor='w')
        ttk.Entry(options_frame, textvariable=self.size_limit, width=10).pack(anchor='w')

        # Progress
        self.progress = ttk.Progressbar(main_frame, orient='horizontal', length=100, mode='determinate')
        self.progress.grid(row=4, column=0, columnspan=3, sticky=('w', 'e'), pady=10)

        # Log
        ttk.Label(main_frame, text="Activity Log:").grid(row=5, column=0, sticky='w')
        self.log_text = Text(main_frame, height=10, width=70)
        self.log_text.grid(row=6, column=0, columnspan=3, sticky=('w', 'e'))
        scrollbar = ttk.Scrollbar(main_frame, orient=VERTICAL, command=self.log_text.yview)
        scrollbar.grid(row=6, column=3, sticky=('n', 's'))
        self.log_text['yscrollcommand'] = scrollbar.set

        # Action buttons
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=7, column=0, columnspan=3, pady=10)
        ttk.Button(button_frame, text="Clean Files", command=self.clean_files).pack(side='left', padx=5)
        ttk.Button(button_frame, text="Clear Log", command=self.clear_log).pack(side='left', padx=5)
        ttk.Button(button_frame, text="Exit", command=self.root.quit).pack(side='right', padx=5)

        main_frame.columnconfigure(1, weight=1)
        main_frame.rowconfigure(6, weight=1)
        self.toggle_output()

    def toggle_output(self):
        self.output_entry.config(state='disabled' if self.same_as_input_var.get() else 'normal')

    def browse_files(self):
        if self.recursive.get():
            folder = filedialog.askdirectory(title="Select folder to clean")
            if folder:
                self.file_entry.delete(0, END)
                self.file_entry.insert(0, folder)
        else:
            files = filedialog.askopenfilenames(
                title="Select files to clean",
                filetypes=[
                    ("All files", "*.*"),
                    ("Images", "*.jpg *.jpeg *.png *.tiff *.bmp *.webp *.gif *.svg *.heic *.cr2 *.nef"),
                    ("Documents", "*.docx *.xlsx *.pdf *.txt *.csv *.odt *.rtf"),
                    ("Media", "*.mp3 *.mp4 *.avi *.wav *.flac *.mkv *.mov"),
                    ("Presentations", "*.pptx *.odp"),
                    ("Archives", "*.zip *.rar *.7z")
                ])
            if files:
                self.file_entry.delete(0, END)
                self.file_entry.insert(0, "; ".join(files))

    def browse_output(self):
        folder = filedialog.askdirectory(title="Select output folder")
        if folder:
            self.output_entry.delete(0, END)
            self.output_entry.insert(0, folder)

    def log(self, message, level='info'):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(END, f"[{timestamp}] {message}\n")
        self.log_text.see(END)
        self.root.update_idletasks()
        if level == 'info':
            self.logger.info(message)
        elif level == 'warning':
            self.logger.warning(message)
        elif level == 'error':
            self.logger.error(message)

    def clear_log(self):
        self.log_text.delete(1.0, END)

    def get_output_path(self, original_path):
        output_dir = os.path.dirname(original_path) if self.same_as_input_var.get() else self.output_entry.get()
        if not output_dir:
            output_dir = os.path.dirname(original_path)
        filename = os.path.basename(original_path)
        name, ext = os.path.splitext(filename)
        new_filename = f"{name}_cleaned{ext}"
        return os.path.join(output_dir, new_filename)

    def cleanup_temp(self):
        temp_dir = tempfile.gettempdir()
        for f in os.listdir(temp_dir):
            if f.startswith('metastripper_'):
                try:
                    os.remove(os.path.join(temp_dir, f))
                except:
                    pass

    def is_valid_zip(self, filepath):
        try:
            with zipfile.ZipFile(filepath, 'r') as zf:
                zf.testzip()
            return True
        except zipfile.BadZipFile:
            return False

    def clean_files(self):
        input_path = self.file_entry.get()
        if not input_path:
            messagebox.showwarning("Warning", "Please select at least one file or folder")
            return

        files = []
        if self.recursive.get():
            if not os.path.isdir(input_path):
                messagebox.showwarning("Warning", "Selected path is not a folder")
                return
            for root, _, filenames in os.walk(input_path):
                files.extend(os.path.join(root, f) for f in filenames)
        else:
            files = [f.strip() for f in input_path.split(";") if f.strip()]

        if not files:
            messagebox.showwarning("Warning", "No files to process")
            return

        total_files = len(files)
        self.progress["maximum"] = total_files
        self.progress["value"] = 0

        try:
            for i, filepath in enumerate(files):
                if not os.path.exists(filepath):
                    self.log(f"File not found: {filepath}", level='warning')
                    continue

                if self.size_limit.get() > 0:
                    size_mb = os.path.getsize(filepath) / (1024 * 1024)
                    if size_mb > self.size_limit.get():
                        self.log(f"Skipping {filepath}: Size {size_mb:.2f} MB exceeds limit", level='warning')
                        continue

                try:
                    self.log(f"Processing: {os.path.basename(filepath)}")
                    output_path = self.get_output_path(filepath)
                    os.makedirs(os.path.dirname(output_path), exist_ok=True)

                    if self.backup.get():
                        backup_path = f"{filepath}.bak"
                        shutil.copy2(filepath, backup_path)
                        self.log(f"Created backup: {backup_path}")

                    ext = os.path.splitext(filepath)[1].lower()
                    if ext in ('.jpg', '.jpeg', '.png', '.tiff', '.bmp', '.webp', '.gif', '.heic', '.cr2', '.nef'):
                        self.clean_image(filepath, output_path)
                    elif ext == '.pdf':
                        self.clean_pdf(filepath, output_path)
                    elif ext == '.docx':
                        self.clean_docx(filepath, output_path)
                    elif ext == '.pptx':
                        self.clean_pptx(filepath, output_path)
                    elif ext in ('.xlsx', '.xls'):
                        self.clean_excel(filepath, output_path)
                    elif ext in ('.odt', '.odp'):
                        self.clean_odf(filepath, output_path)
                    elif ext == '.rtf':
                        self.clean_rtf(filepath, output_path)
                    elif ext in ('.txt', '.csv', '.html'):
                        self.clean_text(filepath, output_path)
                    elif ext in ('.mp3', '.wav', '.flac'):
                        self.clean_audio(filepath, output_path)
                    elif ext in ('.mp4', '.avi', '.mkv', '.mov'):
                        self.clean_video(filepath, output_path)
                    elif ext in ('.zip', '.rar', '.7z'):
                        self.clean_archive(filepath, output_path)
                    else:
                        self.clean_generic(filepath, output_path)

                    self.log(f"Successfully cleaned: {os.path.basename(filepath)}")
                    self.log(f"Saved to: {output_path}")

                except Exception as e:
                    self.log(f"Error processing {os.path.basename(filepath)}: {str(e)}", level='error')

                self.progress["value"] = i + 1
                self.root.update_idletasks()

            messagebox.showinfo("Complete", "Metadata cleaning process finished!")
        finally:
            self.cleanup_temp()

    def clean_image(self, input_path, output_path):
        try:
            ext = os.path.splitext(input_path)[1].lower()
            if ext in ('.heic', '.cr2', '.nef'):
                if not IMAGEIO_AVAILABLE:
                    self.log(f"imageio not installed, copying {input_path} without cleaning", level='warning')
                    shutil.copy2(input_path, output_path)
                    return
                img = imageio.imread(input_path)
                imageio.imwrite(output_path, img)
            elif ext == '.svg':
                shutil.copy2(input_path, output_path)
            else:
                img = Image.open(input_path)
                data = list(img.getdata())
                mode = img.mode
                size = img.size

                new_img = Image.new(mode, size)
                new_img.putdata(data)
                if img.palette:
                    new_img.putpalette(img.getpalette())
                if img.info.get('transparency'):
                    new_img.info['transparency'] = img.info['transparency']

                save_params = {
                    '.png': {'format': 'PNG', 'compress_level': 9},
                    '.jpg': {'format': 'JPEG', 'quality': 95, 'optimize': True},
                    '.jpeg': {'format': 'JPEG', 'quality': 95, 'optimize': True},
                    '.gif': {'format': 'GIF'},
                    '.tiff': {'format': 'TIFF'},
                    '.bmp': {'format': 'BMP'},
                    '.webp': {'format': 'WEBP', 'quality': 95}
                }
                new_img.save(output_path, **save_params.get(ext, {}))
                img.close()
        except Exception as e:
            raise Exception(f"Image cleaning failed: {str(e)}")

    def clean_pdf(self, input_path, output_path):
        try:
            with open(input_path, 'rb') as infile:
                reader = PdfReader(infile)
                if reader.is_encrypted:
                    self.log(f"Encrypted PDF detected: {input_path}, copying without cleaning", level='warning')
                    shutil.copy2(input_path, output_path)
                    return
                writer = PdfWriter()
                for page in reader.pages:
                    writer.add_page(page)
                writer.add_metadata({})
                with open(output_path, "wb") as outfile:
                    writer.write(outfile)
        except Exception as e:
            raise Exception(f"PDF cleaning failed: {str(e)}")

    def clean_docx(self, input_path, output_path):
        try:
            doc = Document(input_path)
            cp = doc.core_properties
            cp.author = cp.title = cp.subject = cp.keywords = cp.comments = cp.last_modified_by = ""
            cp.revision = 1
            cp.category = cp.content_status = cp.identifier = cp.language = cp.version = ""

            if not self.keep_date.get():
                cp.created = cp.modified = cp.last_printed = datetime(2000, 1, 1)
            else:
                cp.modified = datetime.now()

            temp_path = os.path.join(tempfile.gettempdir(), f'metastripper_{os.path.basename(output_path)}')
            doc.save(temp_path)
            shutil.move(temp_path, output_path)
        except Exception as e:
            raise Exception(f"DOCX cleaning failed: {str(e)}")

    def clean_pptx(self, input_path, output_path):
        try:
            temp_path = os.path.join(tempfile.gettempdir(), f'metastripper_{os.path.basename(input_path)}')
            shutil.copy2(input_path, temp_path)
            prs = Presentation(temp_path)
            prs.core_properties.author = prs.core_properties.title = prs.core_properties.subject = ""
            prs.core_properties.keywords = prs.core_properties.comments = ""

            if not self.keep_copyright.get():
                prs.core_properties.category = prs.core_properties.content_status = ""
            if not self.keep_date.get():
                prs.core_properties.created = prs.core_properties.modified = datetime(2000, 1, 1)
            else:
                prs.core_properties.modified = datetime.now()

            temp_output = os.path.join(tempfile.gettempdir(), f'metastripper_{os.path.basename(output_path)}')
            prs.save(temp_output)
            shutil.move(temp_output, output_path)
            os.remove(temp_path)
        except Exception as e:
            raise Exception(f"PPTX cleaning failed: {str(e)}")

    def clean_excel(self, input_path, output_path):
        try:
            if not self.is_valid_zip(input_path):
                self.log(f"Invalid or corrupted Excel file: {input_path}, copying without cleaning", level='warning')
                shutil.copy2(input_path, output_path)
                return
            wb = load_workbook(input_path)
            props = wb.properties
            props.creator = props.title = props.subject = props.keywords = props.description = ""
            props.lastModifiedBy = props.category = props.version = ""

            if not self.keep_date.get():
                props.created = props.modified = datetime(2000, 1, 1)
            wb.save(output_path)
        except Exception as e:
            raise Exception(f"Excel cleaning failed: {str(e)}")

    def clean_odf(self, input_path, output_path):
        try:
            doc = load_odf(input_path)
            meta = doc.getElementsByType(text.Meta)
            for m in meta:
                doc.removeChild(m)
            doc.save(output_path)
        except Exception as e:
            raise Exception(f"ODF cleaning failed: {str(e)}")

    def clean_rtf(self, input_path, output_path):
        try:
            shutil.copy2(input_path, output_path)
        except Exception as e:
            raise Exception(f"RTF cleaning failed: {str(e)}")

    def clean_text(self, input_path, output_path):
        try:
            shutil.copy2(input_path, output_path)
        except Exception as e:
            raise Exception(f"Text file handling failed: {str(e)}")

    def clean_audio(self, input_path, output_path):
        try:
            audio = mutagen.File(input_path)
            if audio:
                audio.delete()
                audio.save()
            shutil.copy2(input_path, output_path)
        except Exception as e:
            raise Exception(f"Audio cleaning failed: {str(e)}")

    def clean_video(self, input_path, output_path):
        try:
            if FFMPEG_AVAILABLE:
                # Check for local ffmpeg.exe when running as executable
                ffmpeg_path = 'ffmpeg'  # Default system ffmpeg
                if getattr(sys, 'frozen', False):  # Running as PyInstaller executable
                    base_path = os.path.dirname(sys.executable)
                    local_ffmpeg = os.path.join(base_path, 'ffmpeg.exe' if sys.platform == 'win32' else 'ffmpeg')
                    if os.path.exists(local_ffmpeg):
                        ffmpeg_path = local_ffmpeg
                stream = ffmpeg.input(input_path)
                stream = ffmpeg.output(stream, output_path, c='copy', map_metadata=-1)
                ffmpeg.run(stream, cmd=ffmpeg_path)
            else:
                self.log(f"ffmpeg-python not installed, copying {input_path} without cleaning", level='warning')
                shutil.copy2(input_path, output_path)
        except Exception as e:
            self.log(f"Video cleaning failed, copying: {str(e)}", level='warning')
            shutil.copy2(input_path, output_path)

    def clean_archive(self, input_path, output_path):
        try:
            shutil.copy2(input_path, output_path)
        except Exception as e:
            raise Exception(f"Archive cleaning failed: {str(e)}")

    def clean_generic(self, input_path, output_path):
        try:
            parser = createParser(input_path)
            if not parser:
                self.log(f"No parser available for {os.path.basename(input_path)} - simple copy")
                shutil.copy2(input_path, output_path)
                return
            with parser:
                metadata = extractMetadata(parser)
                if metadata:
                    with open(input_path, "rb") as src, open(output_path, "wb") as dest:
                        shutil.copyfileobj(src, dest)
                else:
                    shutil.copy2(input_path, output_path)
        except Exception as e:
            raise Exception(f"Generic cleaning failed: {str(e)}")

if __name__ == "__main__":
    root = Tk()
    app = MetaStripper(root)
    root.mainloop()